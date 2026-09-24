"""
Read a banner plan laid out vertically, one row per banner column.

This is the second of the two shapes seen in practice. Where the banner
structure sheet runs horizontally with one column per banner point, a
banner plan runs downwards under named headings:

    Column | Variable | Group | Label | Response | Condition | N | ...
      1    | Total    |   1   |       | Total    | All resp. | 135
      2    | S0       |   2   | Geog. | US       | S0=1      |  75
      3    |          |       |       | EUR      | S0=2,3... |  60

One sheet can hold several banners, each introduced by a title row such as
"Banner 2: US (S0=1)" followed by its own heading row. Group headings come
from the Label column, which is merged down the rows it covers.

Conditions are written in plan syntax, not WinCross syntax, so they are put
through the translator and their status is carried on each point.
"""

import re

import openpyxl

from .logic_translate import translate

DEFAULT_WIDTH = 10

# Heading names, lowercased, mapped to the field they populate
HEADINGS = {
    "column": "column", "col": "column", "#": "column",
    "variable": "variable", "var": "variable",
    "group": "group_no", "grp": "group_no",
    "label": "group", "group label": "group", "heading": "group",
    "response": "label", "banner point": "label", "text": "label",
    "condition": "condition", "logic": "condition", "definition": "condition",
    "n": "n", "base": "n", "base size": "n",
}

TITLE = re.compile(r"^\s*banner\s*(\d+)?\s*[:\-]?\s*(.*)$", re.I)


def looks_like_plan(ws):
    """True when the sheet carries a banner-plan heading row."""
    for r in range(1, min(ws.max_row, 40) + 1):
        seen = {
            str(ws.cell(r, c).value).strip().lower()
            for c in range(1, min(ws.max_column, 15) + 1)
            if ws.cell(r, c).value not in (None, "")
        }
        if {"condition", "response"} <= seen or {"condition", "label"} <= seen:
            return True
    return False


def _resolved(ws, row, col):
    """Cell value, following a merged range to its top-left anchor."""
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            v = ws.cell(rng.min_row, rng.min_col).value
            return "" if v is None else str(v).strip()
    v = ws.cell(row, col).value
    return "" if v is None else str(v).strip()


def _heading_row(ws, row):
    """Map column index -> field name, if `row` is a heading row."""
    mapping = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=row, column=c).value
        if v in (None, ""):
            continue
        key = str(v).strip().lower()
        if key in HEADINGS:
            mapping[c] = HEADINGS[key]
    return mapping if {"condition"} <= set(mapping.values()) else None


def read_plan(path_or_buffer, sheet=None, default_width=DEFAULT_WIDTH,
              lo=0, hi=9999):
    """Return {banner_name: [points]} plus a translation report."""
    wb = openpyxl.load_workbook(path_or_buffer, data_only=True)
    names = [sheet] if sheet else wb.sheetnames
    banners, report = {}, []

    for name in names:
        ws = wb[name]
        if not looks_like_plan(ws):
            continue

        current, cols, title = None, None, None
        for r in range(1, ws.max_row + 1):
            first = _resolved(ws, r, 1)

            heading = _heading_row(ws, r)
            if heading:
                cols = heading
                if current is None:
                    title = title or f"{name}"
                    banners.setdefault(title, [])
                    current = title
                continue

            # A title row: text in column 1, nothing that looks like data
            if first and not cols:
                m = TITLE.match(first)
                if m:
                    title = first.strip()
                continue
            if first and cols and not re.fullmatch(r"\d+", first):
                m = TITLE.match(first)
                if m and "banner" in first.lower():
                    title = first.strip()
                    cols = None
                    current = None
                    continue

            if not cols:
                continue

            row_vals = {field: _resolved(ws, r, c) for c, field in cols.items()}
            if not row_vals.get("condition") and not row_vals.get("label"):
                continue

            if current is None:
                title = title or name
                banners.setdefault(title, [])
                current = title

            t = translate(row_vals.get("condition", ""), lo=lo, hi=hi)
            n = len(banners[current]) + 1
            point = {
                "column": n,
                "label": row_vals.get("label", ""),
                "logic": t.expr,
                "width": default_width,
                "tiers": [row_vals.get("group", "")],
                "super": row_vals.get("group", ""),
                "group": "",
                "variable": row_vals.get("variable", ""),
                "base_n": row_vals.get("n", ""),
                "condition": t.source,
                "status": t.status,
                "note": t.note,
            }
            banners[current].append(point)
            report.append({
                "banner": current, "column": n,
                "label": point["label"], "condition": t.source,
                "expression": t.expr, "status": t.status, "note": t.note,
            })

    if not banners:
        raise ValueError("no banner-plan tables found in this workbook")
    return banners, report


def summarise(report):
    counts = {"ok": 0, "assumed": 0, "blocked": 0}
    for row in report:
        counts[row["status"]] = counts.get(row["status"], 0) + 1
    return counts
