"""
Read a banner structure / banner spec workbook into a list of banner points.

The sheets used in practice share one grid layout:

      col A      | col B      | col C      | ...
    ------------------------------------------------
    (blank)      |                                     <- row 1, usually empty
    (blank)      | super-header, merged across columns  <- row 2
    (blank)      | group heading, merged across columns <- row 3
    (blank)      | column label                         <- row 4
    (blank)      | (sometimes blank)                    <- row 5
    (blank)      | banner logic                         <- last used row

Merged cell ranges carry the spans directly, so headings are read rather
than inferred. The logic row is located as the last row containing data,
which handles sheets with and without the blank spacer row.
"""

import openpyxl

DEFAULT_WIDTH = 10


class SheetError(Exception):
    pass


def _merge_map(ws, row):
    """Map column index -> the value governing it on `row`, following merges."""
    governing = {}
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row:
            value = ws.cell(rng.min_row, rng.min_col).value
            for col in range(rng.min_col, rng.max_col + 1):
                governing[col] = value
    return governing


def _cell(ws, row, col, merges):
    value = merges.get(col, ws.cell(row, col).value)
    if value is None:
        return ""
    return str(value).strip()


def find_rows(ws):
    """Locate the logic row, the label row, and any heading rows above them.

    The logic row is the last populated row; the label row is the one above
    it. Everything populated above that is a heading tier, outermost first.
    Banners vary in how many tiers they carry - some have a super-header and
    a group heading, some only one heading, some none - so the count is read
    from the sheet rather than assumed.
    """
    populated = [
        r for r in range(1, ws.max_row + 1)
        if any(ws.cell(r, c).value not in (None, "") for c in range(1, ws.max_column + 1))
    ]
    if len(populated) < 2:
        raise SheetError(
            f"expected at least 2 populated rows (labels and logic); "
            f"found {len(populated)}")
    logic_row = populated[-1]
    label_row = populated[-2]
    tier_rows = populated[:-2]
    return tier_rows, label_row, logic_row


def read_points(path_or_buffer, sheet=None, default_width=DEFAULT_WIDTH,
                width_overrides=None):
    """Return (points, meta). Each point: super, group, label, logic, width."""
    wb = openpyxl.load_workbook(path_or_buffer, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]

    tier_rows, label_row, logic_row = find_rows(ws)
    merges = [_merge_map(ws, r) for r in tier_rows]

    width_overrides = width_overrides or {}
    points = []
    for col in range(1, ws.max_column + 1):
        label = _cell(ws, label_row, col, {})
        logic = _cell(ws, logic_row, col, {})
        if not label and not logic:
            continue                       # blank stub column on the left
        n = len(points) + 1
        point = {
            "label": label,
            "logic": " ".join(logic.split()),   # normalise internal whitespace
            "width": int(width_overrides.get(n, default_width)),
            "column": n,
            "tiers": [_cell(ws, r, col, m)
                      for r, m in zip(tier_rows, merges)],
        }
        # keep the two-tier names available for display and older callers
        point["super"] = point["tiers"][0] if len(point["tiers"]) > 0 else ""
        point["group"] = point["tiers"][1] if len(point["tiers"]) > 1 else ""
        points.append(point)

    if not points:
        raise SheetError("no banner columns found - check the sheet layout")

    meta = {
        "sheet": ws.title,
        "rows": {"headings": tier_rows, "label": label_row, "logic": logic_row},
        "tiers": len(tier_rows),
        "columns": len(points),
    }
    return points, meta


def parse_width_overrides(text):
    """Parse '1:20, 2:20, 6:20' into {1: 20, 2: 20, 6: 20}."""
    out = {}
    for chunk in (text or "").replace(";", ",").split(","):
        chunk = chunk.strip()
        if not chunk:
            continue
        if ":" not in chunk:
            raise ValueError(f"expected 'column:width', got {chunk!r}")
        col, width = chunk.split(":", 1)
        out[int(col.strip())] = int(width.strip())
    return out
